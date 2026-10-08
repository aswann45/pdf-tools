"""Word compatibility and listener shim checks."""

from __future__ import annotations

import os
import shutil
import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from docx import Document
from pdf_oxide import PdfDocument

from pdf_tools import _pdf_backend, unoserver_listener
from pdf_tools.convert import service
from pdf_tools.models.files import File


@pytest.mark.parametrize("extension", ["doc", "docx"])
def test_word_output_and_bookmark(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    extension: str,
) -> None:
    """Word dispatch keeps output paths and bookmark metadata."""
    source = tmp_path / f"source.{extension}"
    source.touch()
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    calls: list[tuple[Path, Path]] = []

    def convert(input_path: Path, output_path: Path) -> None:
        calls.append((input_path, output_path))
        output_path.touch()

    monkeypatch.setattr(service, "convert_word_with_libreoffice", convert)
    result = service.convert_word_to_pdf(
        File(path=source, bookmark_name="Chapter"),
        output_path=output_dir,
    )

    assert calls == [(source, output_dir / "source.pdf")]
    assert result.path == output_dir / "source.pdf"
    assert result.bookmark_name == "Chapter"


def test_word_output_validation(tmp_path: Path) -> None:
    """Word conversion preserves preflight path errors."""
    source = tmp_path / "source.docx"
    source.touch()
    existing = tmp_path / "existing.pdf"
    existing.touch()
    directory = tmp_path / "directory.pdf"
    directory.mkdir()

    with pytest.raises(FileExistsError):
        service.convert_word_to_pdf(source, existing)
    with pytest.raises(ValueError, match="is a directory"):
        service.convert_word_to_pdf(source, directory)
    with pytest.raises(FileNotFoundError):
        service.convert_word_to_pdf(source, tmp_path / "absent" / "out.pdf")


def test_listener_shim_warns_without_executable() -> None:
    """Existing context-manager callers continue without a listener."""
    with pytest.warns(DeprecationWarning, match="deprecated"):
        with unoserver_listener(port=1234, soffice_path=Path("absent")):
            pass


def test_word_converter_requires_libreoffice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A missing external converter is reported as RuntimeError."""
    monkeypatch.setattr(_pdf_backend.shutil, "which", lambda _: None)
    with pytest.raises(RuntimeError, match="LibreOffice is required"):
        _pdf_backend.convert_word_with_libreoffice(
            tmp_path / "source.doc", tmp_path / "output.pdf"
        )


@pytest.mark.slow
@pytest.mark.skipif(
    os.environ.get("PDF_TOOLS_TEST_LIBREOFFICE") != "1"
    or shutil.which("soffice") is None,
    reason="Set PDF_TOOLS_TEST_LIBREOFFICE=1 to run external conversion",
)
def test_direct_libreoffice_docx(tmp_path: Path) -> None:
    """Direct LibreOffice conversion works with a real DOCX fixture."""
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Direct Word conversion")
    document.save(str(source))

    result = service.convert_word_to_pdf(source)

    pdf = PdfDocument(str(result.path))
    assert pdf.page_count() == 1
    assert "Direct Word conversion" in pdf.extract_text(0)


@pytest.mark.slow
@pytest.mark.skipif(
    os.environ.get("PDF_TOOLS_TEST_LIBREOFFICE") != "1"
    or shutil.which("soffice") is None,
    reason="Set PDF_TOOLS_TEST_LIBREOFFICE=1 to run external conversion",
)
def test_direct_libreoffice_doc(tmp_path: Path) -> None:
    """Legacy binary DOC remains convertible without a listener."""
    source_docx = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Legacy Word conversion")
    document.save(str(source_docx))

    with TemporaryDirectory() as directory:
        profile = (Path(directory) / "profile").as_uri()
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "doc",
                "--outdir",
                str(tmp_path),
                str(source_docx),
            ],
            check=True,
            capture_output=True,
            timeout=120,
        )

    result = service.convert_word_to_pdf(tmp_path / "source.doc")
    pdf = PdfDocument(str(result.path))
    assert "Legacy Word conversion" in pdf.extract_text(0)
